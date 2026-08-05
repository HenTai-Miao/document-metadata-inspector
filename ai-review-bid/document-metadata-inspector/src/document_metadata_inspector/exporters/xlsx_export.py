"""Truncation-free, formula-safe XLSX report projection."""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from typing import TYPE_CHECKING, Final, cast

from openpyxl import Workbook

from document_metadata_inspector.models import ComparisonReport

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.cell.cell import Cell
    from openpyxl.worksheet.worksheet import Worksheet

    from document_metadata_inspector.exporters.contracts import ExportableReport
    from document_metadata_inspector.models import (
        DocumentReport,
        EmbeddedDocumentReport,
        EvidenceObservation,
        JsonValue,
        ParserRun,
    )

type CellValue = str | int | float | bool | None

_CELL_TEXT_LIMIT: Final = 30_000
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FORENSIC_TOKENS: Final = (
    "annotation",
    "comment",
    "revision",
    "signature",
    "embedded",
    "attachment",
    "external_relationship",
    "active_content",
    "encryption",
    "permissions",
    "trailer",
    "acroform",
    "macro",
    "activex",
)


def _text(value: object) -> str:
    text = str(value)
    return _ILLEGAL_XML.sub(lambda match: f"\\u{ord(match.group()):04x}", text)


def _json(value: JsonValue) -> str:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )


def _chunks(value: str) -> tuple[str, ...]:
    if not value:
        return ("",)
    return tuple(
        value[index : index + _CELL_TEXT_LIMIT]
        for index in range(0, len(value), _CELL_TEXT_LIMIT)
    )


def _append(worksheet: Worksheet, values: tuple[CellValue, ...]) -> None:
    first_cell = worksheet.cell(row=1, column=1)
    row = (
        1
        if worksheet.max_row == 1 and first_cell.value is None
        else worksheet.max_row + 1
    )
    for column, value in enumerate(values, start=1):
        cell = cast("Cell", worksheet.cell(row=row, column=column))
        if isinstance(value, str):
            cell.value = _text(value)
            cell.data_type = "s"
        else:
            cell.value = value


def _sheet(
    workbook: Workbook,
    title: str,
    headers: tuple[str, ...],
) -> Worksheet:
    worksheet = cast("Worksheet", workbook.create_sheet(title))
    _append(worksheet, headers)
    worksheet.freeze_panes = "A2"
    return worksheet


def _observation_rows(
    document: DocumentReport,
    index: int,
    observation: EvidenceObservation,
) -> tuple[tuple[CellValue, ...], ...]:
    raw_parts = _chunks(_json(observation.raw_value))
    normalized_parts = _chunks(_json(observation.normalized_value))
    part_count = max(len(raw_parts), len(normalized_parts))
    return tuple(
        (
            document.display_name,
            index,
            observation.canonical_name,
            observation.raw_name,
            part_index + 1,
            part_count,
            raw_parts[part_index] if part_index < len(raw_parts) else "",
            normalized_parts[part_index] if part_index < len(normalized_parts) else "",
            observation.value_type,
            observation.source_kind,
            observation.location,
            observation.parser_name,
            observation.parser_version,
            observation.status.value,
            " | ".join(observation.warnings),
        )
        for part_index in range(part_count)
    )


def _is_forensic(observation: EvidenceObservation) -> bool:
    searchable = f"{observation.canonical_name} {observation.source_kind}".casefold()
    return any(token in searchable for token in _FORENSIC_TOKENS)


def _totals(report: ExportableReport) -> dict[str, int]:
    documents = report.documents
    return {
        "document_count": len(documents),
        "observation_count": sum(
            len(item.observations)
            + sum(len(child.observations) for child in item.embedded_documents)
            for item in documents
        ),
        "embedded_document_count": sum(
            len(item.embedded_documents) for item in documents
        ),
        "parser_run_count": sum(len(item.parser_runs) for item in documents),
        "match_count": len(report.matches)
        if isinstance(report, ComparisonReport)
        else 0,
        "warning_count": len(report.warnings)
        + sum(len(item.warnings) for item in documents)
        + sum(len(run.warnings) for item in documents for run in item.parser_runs)
        + sum(
            len(child.warnings)
            for item in documents
            for child in item.embedded_documents
        ),
    }


@dataclass(frozen=True, slots=True)
class _Sheets:
    summary: Worksheet
    files: Worksheet
    canonical: Worksheet
    matches: Worksheet
    raw: Worksheet
    forensic: Worksheet
    embedded_documents: Worksheet
    embedded_metadata: Worksheet
    coverage: Worksheet
    warnings: Worksheet


def _create_sheets(workbook: Workbook) -> _Sheets:
    return _Sheets(
        summary=_sheet(workbook, "Summary", ("metric", "value")),
        files=_sheet(
            workbook,
            "Files",
            (
                "display_name",
                "source_path_sha256",
                "sha256",
                "size_bytes",
                "declared_format",
                "family",
                "expected_container",
                "observed_container",
                "identity_status",
                "document_status",
                "encrypted",
                "password_required",
                "observation_count",
                "parser_run_count",
                "warning_count",
            ),
        ),
        canonical=_sheet(
            workbook,
            "Canonical Metadata",
            (
                "document",
                "canonical_name",
                "part_index",
                "part_count",
                "normalized_value_part",
                "raw_value_part",
                "status",
                "source_kind",
                "location",
                "parser",
            ),
        ),
        matches=_sheet(
            workbook,
            "Identity Matches",
            (
                "match_index",
                "trace_kind",
                "canonical_name",
                "part_index",
                "part_count",
                "normalized_value_part",
                "documents",
                "document_ids",
                "field_names",
                "evidence_count",
                "raw_values_part",
            ),
        ),
        raw=_sheet(
            workbook,
            "Raw Observations",
            (
                "document",
                "observation_index",
                "canonical_name",
                "raw_name",
                "part_index",
                "part_count",
                "raw_value_part",
                "normalized_value_part",
                "value_type",
                "source_kind",
                "location",
                "parser_name",
                "parser_version",
                "status",
                "warnings",
            ),
        ),
        forensic=_sheet(
            workbook,
            "Forensic Evidence",
            (
                "document",
                "canonical_name",
                "raw_name",
                "part_index",
                "part_count",
                "raw_value_part",
                "source_kind",
                "location",
                "parser",
            ),
        ),
        embedded_documents=_sheet(
            workbook,
            "Embedded Documents",
            (
                "document",
                "resource_id",
                "depth",
                "path",
                "content_type",
                "size_bytes",
                "sha256",
                "parent_id",
                "parsers",
                "observation_count",
                "warnings",
            ),
        ),
        embedded_metadata=_sheet(
            workbook,
            "Embedded Metadata",
            (
                "document",
                "resource_id",
                "path",
                "observation_index",
                "canonical_name",
                "raw_name",
                "raw_value_part",
                "normalized_value_part",
                "part_index",
                "part_count",
                "value_type",
                "source_kind",
                "location",
                "parser_name",
                "parser_version",
                "status",
                "warnings",
            ),
        ),
        coverage=_sheet(
            workbook,
            "Parser Coverage",
            (
                "document",
                "parser_name",
                "parser_version",
                "status",
                "observation_count",
                "duration_ms",
                "error_code",
                "warnings",
            ),
        ),
        warnings=_sheet(
            workbook,
            "Warnings Errors",
            ("scope", "document", "parser", "kind", "value"),
        ),
    )


def _write_observation(
    sheets: _Sheets,
    document: DocumentReport,
    index: int,
    observation: EvidenceObservation,
) -> None:
    for row in _observation_rows(document, index, observation):
        _append(sheets.raw, row)
        _append(
            sheets.canonical,
            (
                row[0],
                row[2],
                row[4],
                row[5],
                row[7],
                row[6],
                row[13],
                row[9],
                row[10],
                row[11],
            ),
        )
    if not _is_forensic(observation):
        return
    raw_parts = _chunks(_json(observation.raw_value))
    for part_index, raw_part in enumerate(raw_parts, start=1):
        _append(
            sheets.forensic,
            (
                document.display_name,
                observation.canonical_name,
                observation.raw_name,
                part_index,
                len(raw_parts),
                raw_part,
                observation.source_kind,
                observation.location,
                observation.parser_name,
            ),
        )


def _write_parser_run(
    sheets: _Sheets,
    document: DocumentReport,
    parser_run: ParserRun,
) -> None:
    _append(
        sheets.coverage,
        (
            document.display_name,
            parser_run.parser_name,
            parser_run.parser_version,
            parser_run.status.value,
            len(parser_run.observations)
            + sum(len(child.observations) for child in parser_run.embedded_documents),
            parser_run.duration_ms,
            parser_run.error_code or "",
            " | ".join(parser_run.warnings),
        ),
    )
    if parser_run.error_code:
        _append(
            sheets.warnings,
            (
                "parser",
                document.display_name,
                parser_run.parser_name,
                "error",
                parser_run.error_code,
            ),
        )
    for warning in parser_run.warnings:
        _append(
            sheets.warnings,
            (
                "parser",
                document.display_name,
                parser_run.parser_name,
                "warning",
                warning,
            ),
        )


def _write_embedded_document(
    sheets: _Sheets,
    document: DocumentReport,
    child: EmbeddedDocumentReport,
) -> None:
    _append(
        sheets.embedded_documents,
        (
            document.display_name,
            child.resource_id,
            child.depth,
            child.path,
            child.content_type,
            child.size_bytes,
            child.sha256,
            child.parent_id,
            " | ".join(child.parser_names),
            len(child.observations),
            " | ".join(child.warnings),
        ),
    )
    for warning in child.warnings:
        _append(
            sheets.warnings,
            ("embedded", document.display_name, "", "warning", warning),
        )
    for index, observation in enumerate(child.observations, start=1):
        raw_parts = _chunks(_json(observation.raw_value))
        normalized_parts = _chunks(_json(observation.normalized_value))
        part_count = max(len(raw_parts), len(normalized_parts))
        for part_index in range(part_count):
            _append(
                sheets.embedded_metadata,
                (
                    document.display_name,
                    child.resource_id,
                    child.path,
                    index,
                    observation.canonical_name,
                    observation.raw_name,
                    raw_parts[part_index] if part_index < len(raw_parts) else "",
                    normalized_parts[part_index]
                    if part_index < len(normalized_parts)
                    else "",
                    part_index + 1,
                    part_count,
                    observation.value_type,
                    observation.source_kind,
                    observation.location,
                    observation.parser_name,
                    observation.parser_version,
                    observation.status.value,
                    " | ".join(observation.warnings),
                ),
            )


def _write_document(sheets: _Sheets, document: DocumentReport) -> None:
    observed = document.identity.observed_container
    _append(
        sheets.files,
        (
            document.display_name,
            document.source_path_sha256 or "",
            document.sha256,
            document.size_bytes,
            document.identity.declared_format.value,
            document.identity.family.value,
            document.identity.expected_container.value,
            observed.value if observed is not None else "",
            document.identity.status.value,
            document.status.value,
            document.encrypted,
            document.password_required,
            len(document.observations)
            + sum(len(child.observations) for child in document.embedded_documents),
            len(document.parser_runs),
            len(document.warnings),
        ),
    )
    for warning in document.warnings:
        _append(
            sheets.warnings,
            ("document", document.display_name, "", "warning", warning),
        )
    for index, observation in enumerate(document.observations, start=1):
        _write_observation(sheets, document, index, observation)
    for child in document.embedded_documents:
        _write_embedded_document(sheets, document, child)
    for parser_run in document.parser_runs:
        _write_parser_run(sheets, document, parser_run)


def _write_matches(sheets: _Sheets, report: ExportableReport) -> None:
    if not isinstance(report, ComparisonReport):
        return
    for index, match in enumerate(report.matches, start=1):
        normalized_parts = _chunks(_json(match.normalized_value))
        raw_parts = _chunks(_json(match.raw_values))
        part_count = max(len(normalized_parts), len(raw_parts))
        for part_index in range(part_count):
            _append(
                sheets.matches,
                (
                    index,
                    match.trace_kind.value,
                    match.canonical_name,
                    part_index + 1,
                    part_count,
                    normalized_parts[part_index]
                    if part_index < len(normalized_parts)
                    else "",
                    " | ".join(match.document_names),
                    " | ".join(match.document_ids),
                    " | ".join(match.field_names),
                    match.evidence_count,
                    raw_parts[part_index] if part_index < len(raw_parts) else "",
                ),
            )


def _write_summary(
    workbook: Workbook,
    sheets: _Sheets,
    report: ExportableReport,
) -> None:
    for metric, value in _totals(report).items():
        _append(sheets.summary, (metric, value))
    for worksheet in workbook.worksheets:
        if worksheet.title != sheets.summary.title:
            _append(
                sheets.summary,
                (f"sheet_rows:{worksheet.title}", max(worksheet.max_row - 1, 0)),
            )


def write_xlsx_report(report: ExportableReport, path: Path) -> None:
    """Write every report row as formula-safe text without cell truncation."""
    workbook = Workbook()
    try:
        default = workbook.active
        if default is not None:
            workbook.remove(default)
        sheets = _create_sheets(workbook)
        for warning in report.warnings:
            _append(sheets.warnings, ("report", "", "", "warning", warning))
        for document in report.documents:
            _write_document(sheets, document)
        _write_matches(sheets, report)
        _write_summary(workbook, sheets, report)
        workbook.save(path)
    finally:
        workbook.close()


__all__ = ("write_xlsx_report",)
