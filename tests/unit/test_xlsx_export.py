from pathlib import Path

from openpyxl import load_workbook

from document_metadata_inspector.exporters import (
    ExportFormat,
    ExportTarget,
    export_report,
)
from tests.export_factories import make_export_report

_LONG_TEXT = "元数据<>&" * 10_000


def test_xlsx_projection_has_all_required_sheets_and_exact_counts(
    tmp_path: Path,
) -> None:
    report = make_export_report(_LONG_TEXT)
    target = tmp_path / "report.xlsx"

    _ = export_report(report, ExportTarget(path=target, format=ExportFormat.XLSX))

    workbook = load_workbook(target, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Summary",
            "Files",
            "Canonical Metadata",
            "Identity Matches",
            "Raw Observations",
            "Forensic Evidence",
            "Embedded Documents",
            "Embedded Metadata",
            "Parser Coverage",
            "Warnings Errors",
        ]
        summary: dict[str, int] = {}
        for metric, value in workbook["Summary"].iter_rows(
            min_row=2,
            values_only=True,
        ):
            if isinstance(metric, str) and isinstance(value, int):
                summary[metric] = value  # noqa: PERF403 - narrows openpyxl unions.
        assert summary["document_count"] == 1
        assert summary["observation_count"] == 3
        assert summary["parser_run_count"] == 1
        assert summary["embedded_document_count"] == 1

        rows = list(workbook["Raw Observations"].iter_rows(min_row=2, values_only=True))
        header = [str(cell.value) for cell in workbook["Raw Observations"][1]]
        raw_index = header.index("raw_value_part")
        name_index = header.index("canonical_name")
        long_parts: list[str] = []
        for row in rows:
            raw_part = row[raw_index]
            if row[name_index] == "author" and isinstance(raw_part, str):
                long_parts.append(raw_part)
        assert "".join(long_parts) == '"' + _LONG_TEXT + '"'

        command_rows = [
            row for row in rows if row[name_index] == "custom_property.command"
        ]
        assert command_rows[0][raw_index] == '"=2+2"'

        embedded_rows = list(
            workbook["Embedded Documents"].iter_rows(min_row=2, values_only=True)
        )
        assert embedded_rows[0][3] == "attachments/child.docx"
        assert embedded_rows[0][6] == "4" * 64

        child_rows = list(
            workbook["Embedded Metadata"].iter_rows(min_row=2, values_only=True)
        )
        assert child_rows[0][4] == "author"
        assert child_rows[0][6] == '"Child <Author>"'
    finally:
        workbook.close()


def test_xlsx_export_reopens_without_formula_cells(tmp_path: Path) -> None:
    target = tmp_path / "report.xlsx"
    _ = export_report(
        make_export_report(),
        ExportTarget(path=target, format=ExportFormat.XLSX),
    )

    workbook = load_workbook(target, read_only=False, data_only=False)
    try:
        formulas = [
            cell.value
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        ]
        assert formulas == []
    finally:
        workbook.close()
